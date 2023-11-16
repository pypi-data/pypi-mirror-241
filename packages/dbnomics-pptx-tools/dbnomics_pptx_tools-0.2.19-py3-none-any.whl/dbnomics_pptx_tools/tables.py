from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from datetime import date
from typing import TypeAlias, cast

import daiquiri
import pandas as pd
from isodate import Duration
from pandas import DataFrame
from parsy import ParseError, any_char, char_from, regex, seq, string
from pptx.table import Table, _Cell, _Row, _RowCollection
from tabulate import tabulate

from dbnomics_pptx_tools.data_loader import ShapeDataLoader
from dbnomics_pptx_tools.formatters import format_number
from dbnomics_pptx_tools.metadata import ColumnsSpec, SeriesSpec, TableSpec
from dbnomics_pptx_tools.repo import SeriesRepo

logger = daiquiri.getLogger(__name__)


@dataclass
class AnnualPeriod:
    year: int

    def __str__(self) -> str:
        return str(self.year)

    def to_pandas_period(self) -> pd.Period:
        return pd.Period(str(self), freq="A")


@dataclass
class MonthlyPeriod:
    year: int
    month: int

    def __str__(self) -> str:
        return f"{self.year}-{self.month:02}"

    def to_pandas_period(self) -> pd.Period:
        return pd.Period(str(self), freq="M")


@dataclass
class QuarterlyPeriod:
    year: int
    quarter: int

    def __str__(self) -> str:
        return f"{self.year}-Q{self.quarter}"

    def to_pandas_period(self) -> pd.Period:
        return pd.Period(str(self), freq="Q")


Period: TypeAlias = AnnualPeriod | MonthlyPeriod | QuarterlyPeriod


@dataclass
class TableZones:
    first_data_row_index: int
    header_row_index: int
    period_count: int
    periods: list[Period] | None


def extract_table_zones(table: Table, *, table_spec: TableSpec | None) -> TableZones | None:
    header_row_index = find_header_row_index(table, table_spec=table_spec)
    if header_row_index is None:
        logger.debug("Could not find the row corresponding to the table header")
        return None

    rows = cast(_RowCollection, table.rows)
    header_cells = list(rows[header_row_index].cells)
    period_cells = header_cells[1:]

    if table_spec is not None and table_spec.columns is None:
        periods = parse_header_periods(period_cells)
        if periods is None:
            logger.debug("Could not parse the periods in the table header")
            return None
    else:
        logger.debug('Skipping parsing periods of table header because "columns" are defined in the table spec')
        periods = None

    first_data_row_index = find_first_data_row_index(table, header_row_index=header_row_index)
    if first_data_row_index is None:
        logger.debug("Could not find the first data row of the table")
        return None

    return TableZones(
        first_data_row_index=first_data_row_index,
        header_row_index=header_row_index,
        period_count=len(period_cells),
        periods=periods,
    )


def find_first_data_row_index(table: Table, *, header_row_index: int) -> int | None:
    rows = list(cast(Iterable[_Row], table.rows))
    start = header_row_index + 1
    for row_index, row in enumerate(rows[start:], start=start):
        first_cell_text = row.cells[0].text.strip()
        if not first_cell_text:
            continue
        return row_index
    return None


def find_header_row_index(table: Table, *, table_spec: TableSpec | None) -> int | None:
    rows = cast(Iterable[_Row], table.rows)
    for row_index, row in enumerate(rows):
        if row.cells[0].is_merge_origin:
            if row_index == 0:
                continue
            raise NotImplementedError
        if table_spec is not None and row.cells[0].text == table_spec.header_first_cell:
            return row_index
    return None


def find_latest_column_period(table_series_df: DataFrame) -> pd.Period:
    logger.debug("Finding the period of the latest column from table series")
    return pd.Period(
        max(sub_df["original_period"].dropna().iloc[-1] for _, sub_df in table_series_df.groupby("series_id"))
    )


def format_table(table: Table) -> str:
    rows = cast(Iterable[_Row], table.rows)
    tabular_data = [[cell.text for cell in row.cells] for row in rows]
    return tabulate(tabular_data)


def generate_periods(
    columns_spec: ColumnsSpec, *, table_series_df: DataFrame, table_zones: TableZones
) -> Sequence[Period]:
    frequency = columns_spec.frequency

    end_period_offset = columns_spec._end_period_offset  # noqa: SLF001
    end_period = (
        find_latest_column_period(table_series_df)
        if end_period_offset is None
        else get_end_period_with_offset(end_period_offset)
    )
    logger.debug("The period of the latest column is %r", end_period)

    if frequency == "annual":
        return generate_annual_periods(end_period=end_period, table_zones=table_zones)

    if frequency == "monthly":
        return generate_monthly_periods(end_period=end_period, table_zones=table_zones)

    if frequency == "quarterly":
        return generate_quarterly_periods(end_period=end_period, table_zones=table_zones)

    msg = f"Invalid frequency: {frequency!r}"
    raise ValueError(msg)


def generate_annual_periods(*, end_period: pd.Period, table_zones: TableZones) -> Sequence[AnnualPeriod]:
    periods = pd.period_range(end=end_period, periods=table_zones.period_count, freq="A")
    return [AnnualPeriod(period.year) for period in periods]


def generate_monthly_periods(*, end_period: pd.Period, table_zones: TableZones) -> Sequence[MonthlyPeriod]:
    periods = pd.period_range(end=end_period, periods=table_zones.period_count, freq="M")
    return [MonthlyPeriod(period.year, period.month) for period in periods]


def generate_quarterly_periods(*, end_period: pd.Period, table_zones: TableZones) -> Sequence[QuarterlyPeriod]:
    periods = pd.period_range(end=end_period, periods=table_zones.period_count, freq="Q")
    return [QuarterlyPeriod(period.year, period.quarter) for period in periods]


def get_end_period_with_offset(end_period_offset: Duration) -> pd.Period:
    logger.debug(
        "Using the date of today with an offset of %r to determine the period of the latest column", end_period_offset
    )
    return pd.Period(date.today() + end_period_offset)  # noqa: DTZ011


def iter_table_data_row_index_with_series_spec(
    table: Table, *, table_spec: TableSpec, table_zones: TableZones
) -> Iterator[tuple[int, SeriesSpec]]:
    row_index_by_label = {v: k for k, v in iter_table_data_rows(table, table_zones=table_zones)}
    for series_spec in table_spec.series:
        series_name = series_spec.name
        row_index = row_index_by_label.get(series_name)
        if row_index is None:
            logger.warning("Could not find the row corresponding to the series name %r, ignoring", series_name)
            continue
        yield row_index, series_spec


def iter_table_data_rows(table: Table, *, table_zones: TableZones) -> Iterator[tuple[int, str]]:
    rows = list(cast(Iterable[_Row], table.rows))
    first_data_row_index = table_zones.first_data_row_index

    for row_index, row in enumerate(rows[first_data_row_index:], start=first_data_row_index):
        row_label = row.cells[0].text.strip()
        yield row_index, row_label


def parse_header_period(text: str) -> Period:
    def to_full_year(x: str) -> int:
        return int(x) + 2000

    full_year = regex(r"[0-9]{4}").desc("4 digit year")
    short_year = regex(r"[0-9]{2}").desc("2 digit year")
    annual_period = full_year.map(int)
    quarterly_period = seq(year=short_year.map(to_full_year), quarter=(string("Q") >> char_from("1234").map(int)))
    period = annual_period.map(AnnualPeriod) | quarterly_period.combine_dict(QuarterlyPeriod)
    attribute = string("(") >> any_char << string(")")
    period_header = period << attribute.optional()
    return cast(Period, period_header.parse(text))


def parse_header_periods(period_cells: list[_Cell]) -> list[Period] | None:
    period_cells_text = [cell.text for cell in period_cells]
    logger.debug("Parsing period cells: %r", period_cells_text)

    has_errors = False

    def iter_parsed() -> Iterator[Period]:
        nonlocal has_errors
        for text_pos, text in enumerate(period_cells_text):
            try:
                yield parse_header_period(text)
            except ParseError:
                logger.exception("Could not parse period at index %d from text %r", text_pos, text)
                has_errors = True

    return None if has_errors else list(iter_parsed())


def replace_cell_text(cell: _Cell, text: str) -> None:
    runs = cell.text_frame.paragraphs[0].runs
    if not runs:
        msg = "No text in cell, do not know which format (font, alignment, etc.) to use"
        raise ValueError(msg)
    runs[0].text = text


def update_table(table: Table, *, repo: SeriesRepo, table_spec: TableSpec, table_zones: TableZones) -> None:
    table_series_df = ShapeDataLoader(repo).load_shape_df(table_spec)

    columns_spec = table_spec.columns
    if columns_spec is None:
        assert table_zones.periods is not None
        periods = table_zones.periods
    else:
        periods = list(generate_periods(columns_spec, table_series_df=table_series_df, table_zones=table_zones))
        update_table_header(table, columns_spec=columns_spec, periods=periods, table_zones=table_zones)

    for row_index, series_spec in iter_table_data_row_index_with_series_spec(
        table, table_spec=table_spec, table_zones=table_zones
    ):
        series_id = series_spec.id
        series_name = series_spec.name
        logger.debug("Processing table row %d named %r related to series %r", row_index, series_name, series_id)

        row_series_df = table_series_df.query("series_id == @series_id")

        for column_index, period in enumerate(periods, start=1):
            update_table_cell(
                table,
                column_index=column_index,
                period=str(period),
                row_index=row_index,
                series_id=series_id,
                row_series_df=row_series_df,
            )


def update_table_cell(
    table: Table, *, column_index: int, period: str, row_index: int, row_series_df: DataFrame, series_id: str
) -> None:
    cell = table.cell(row_index, column_index)
    observations = row_series_df[row_series_df["original_period"] == period].value

    if observations.empty:
        logger.debug(
            "Period %r requested for table, but not found in series %r, keeping current value (%r)",
            period,
            series_id,
            cell.text,
        )
    elif len(observations) > 1:
        logger.warning("Many observations found for period %r in series %r, ignoring period", period, series_id)
    else:
        observation = cast(float, observations.values[0])
        replace_cell_text(cell, format_number(observation))


def update_table_header(
    table: Table, *, columns_spec: ColumnsSpec, periods: list[Period], table_zones: TableZones
) -> None:
    row_index = table_zones.header_row_index
    logger.debug("Updating the table header (row %d) with periods %r", row_index, [str(p) for p in periods])
    for column_index, period in enumerate(periods, start=1):
        cell = table.cell(row_index, column_index)
        pd_period = period.to_pandas_period()
        period_str = pd_period.strftime(columns_spec.period_format)
        replace_cell_text(cell, period_str)
